"""Compare workbook metadata with a frozen manifest; no video reads/decodes."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from data import parse_source_timestamp


def audit(data_dir, split_path):
    from openpyxl import load_workbook
    split = json.loads(Path(split_path).read_text(encoding="utf-8-sig"))
    frozen = {name: row for row in split["source"] for name in row["names"]}
    groups = {row["sha256"]: set() for row in split["source"]}
    book = load_workbook(Path(data_dir) / "open_labels.xlsx", read_only=True, data_only=True)
    try:
        rows = list(book["labels"].iter_rows(values_only=True))[1:]
    finally:
        book.close()
    rows = [row for row in rows if any(value is not None for value in row)]
    seen = set()
    mismatch = 0
    for row in rows:
        name, value, label = (list(row) + [None] * 3)[:3]
        if name in seen or name not in frozen:
            raise ValueError("Workbook membership differs from frozen source")
        seen.add(name)
        timestamp = parse_source_timestamp(value)
        if str(label).strip() not in {"нарушение", "без нарушения"}:
            raise ValueError("Unrecognized label")
        mismatch += (str(label).strip() == "нарушение") != frozen[name]["positive"]
        if timestamp is not None:
            groups[frozen[name]["sha256"]].add(timestamp)
    if seen != set(frozen):
        raise ValueError("Workbook and frozen source memberships differ")
    timestamp_mismatch = sum(groups[row["sha256"]] != set(row["timestamps"])
        for row in split["source"])
    return {"metadata_rows": len(rows), "unique_content_groups": len(groups),
        "label_mismatches": mismatch, "timestamp_union_mismatches": timestamp_mismatch,
        "frozen_metadata_unchanged": mismatch == timestamp_mismatch == 0,
        "video_reads": 0, "video_decodes": 0, "model_loads": 0}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-dir", type=Path, required=True)
    parser.add_argument("--split", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = audit(args.data_dir, args.split)
    if args.output:
        args.output.write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, indent=2))
