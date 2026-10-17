"""Private label/split I/O only; this module never imports a vision model.

Group identifiers are storage/join keys, never predictor features. Existing
frozen splits can be relocated through --data-dir without changing membership.
"""
from __future__ import annotations

import hashlib
import json
import math
from numbers import Real
import os
from pathlib import Path, PureWindowsPath
import re

SEED = 26017


def local_path(path, must_exist=True):
    """Resolve existing paths; accommodate Windows sandbox final-handle denial.

    Some Windows sandboxes allow stat/read but deny GetFinalPathNameByHandle
    on directories. In that case accept lexical normalization only after
    rejecting every symlink/junction ancestor, so containment cannot be bypassed.
    """
    path = Path(os.path.abspath(path))
    try:
        return path.resolve(strict=must_exist)
    except PermissionError as error:
        if os.name != "nt" or getattr(error, "winerror", None) != 5:
            raise
        for part in (path, *path.parents):
            if part.is_symlink() or part.is_junction():
                raise ValueError("Cannot safely resolve a Windows reparse path") from error
        if must_exist and not path.exists():
            raise FileNotFoundError(path)
        return path


def digest(path):
    result = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            result.update(block)
    return result.hexdigest()


def parse_timestamp(value):
    """Finite nonnegative seconds or a complete MM:SS[.fraction] string.

    Missing values return None. Booleans, signed/garbage strings and partial
    matches are rejected rather than stripped or reinterpreted.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("Boolean is not a timestamp")
    if isinstance(value, Real):
        seconds = float(value)
    elif isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        if re.fullmatch(r"[0-9]+(?:\.[0-9]+)?", text):
            seconds = float(text)
        else:
            match = re.fullmatch(r"([0-9]+):([0-9]{2}(?:\.[0-9]+)?)", text)
            if not match or float(match[2]) >= 60:
                raise ValueError("Use complete nonnegative seconds or MM:SS")
            seconds = int(match[1]) * 60 + float(match[2])
    else:
        raise ValueError("Unsupported timestamp value type")
    if not math.isfinite(seconds) or seconds < 0:
        raise ValueError("Timestamp must be finite and nonnegative")
    return seconds


def parse_source_timestamp(value):
    """Adapt the supplied workbook's anchored Russian-unit timestamp schema.

    This explicitly declared schema is separate from the generic seconds/MM:SS
    parser. Whole minutes are optional; fractional seconds use dot or comma.
    No unrelated suffix, sign or partial match is accepted.
    """
    if isinstance(value, str):
        match = re.fullmatch(r"(?:([0-9]+)\s*минут(?:а|ы)?\s*)?"
            r"([0-9]+(?:[.,][0-9]+)?)\s*секунд(?:а|ы)?", value.strip(), flags=re.IGNORECASE)
        if match:
            seconds = parse_timestamp(match[2].replace(",", "."))
            if match[1] is not None:
                if seconds >= 60:
                    raise ValueError("Seconds must be below60 when minutes are present")
                seconds += int(match[1]) * 60
            return parse_timestamp(seconds)
    return parse_timestamp(value)


def read_source(data_dir):
    from openpyxl import load_workbook
    root = local_path(data_dir)
    book = load_workbook(root / "open_labels.xlsx", read_only=True, data_only=True)
    try:
        rows = list(book["labels"].iter_rows(values_only=True))[1:]
    finally:
        book.close()
    rows = [row for row in rows if any(value is not None for value in row)]
    paths = sorted(path for path in (root / "videos").iterdir()
                   if path.is_file() and path.suffix.casefold() == ".mov")
    files = {path.stem: path for path in paths}
    if len(files) != len(paths) or len(files) != len(rows):
        raise ValueError("Workbook and video inventory differ or names collide")
    groups, seen_names = {}, set()
    for row in rows:
        name, raw_timestamp, raw_label = (list(row) + [None] * 3)[:3]
        if name not in files or name in seen_names:
            raise ValueError("Missing or repeated exact Unicode video identifier")
        seen_names.add(name)
        label = str(raw_label).strip()
        if label not in {"нарушение", "без нарушения"}:
            raise ValueError("Unsupported binary source label")
        positive = label == "нарушение"
        timestamp = parse_source_timestamp(raw_timestamp)
        timestamps = [] if timestamp is None else [timestamp]
        if positive and not timestamps:
            raise ValueError("Positive fitting source lacks a sparse timestamp")
        path = files[name]
        sha = digest(path)
        group = groups.setdefault(sha, {"sha256": sha,
            "relative_path": path.relative_to(root).as_posix(),
            "positive": positive, "names": [], "timestamps": []})
        if group["positive"] != positive:
            raise ValueError("Identical video bytes have conflicting binary labels")
        group["names"].append(name)
        group["timestamps"] = sorted(set(group["timestamps"] + timestamps))
    return sorted(groups.values(), key=lambda row: row["sha256"])


def freeze_split(data_dir, destination):
    """Original deterministic class-stratified 20%/20% policy, never overwrite."""
    import numpy as np
    destination = Path(destination)
    if destination.exists():
        raise ValueError("Frozen split already exists; reuse it explicitly")
    rows = read_source(data_dir)
    rng = np.random.default_rng(SEED)
    assignments = {}
    for positive in (False, True):
        members = [row for row in rows if row["positive"] == positive]
        rng.shuffle(members)
        # This is a CV feasibility constraint, not an expected dataset volume.
        n_test = max(2, round(len(members) * .2))
        n_validation = max(2, round(len(members) * .2))
        if len(members) - n_test - n_validation < 4:
            raise ValueError("Each class needs four training groups after holdouts")
        for index, row in enumerate(members):
            assignments[row["sha256"]] = ("test" if index < n_test else
                "validation" if index < n_test + n_validation else "train")
    fingerprint = hashlib.sha256(json.dumps([(row["sha256"], row["positive"],
        sorted(row["timestamps"])) for row in rows]).encode()).hexdigest()
    result = {"schema_version": 1, "seed": SEED, "source_fingerprint": fingerprint,
              "assignments": assignments, "source": rows}
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("x", encoding="utf-8") as stream:
        json.dump(result, stream, indent=2)
    return {"groups": len(rows), "counts": {phase: sum(value == phase for value
        in assignments.values()) for phase in ("train", "validation", "test")},
        "source_decodes": 0, "model_loads": 0}


def resolve_groups(split_path, data_dir, phase):
    if phase not in {"train", "validation"}:
        raise ValueError("No test-processing mode exists")
    split = json.loads(Path(split_path).read_text(encoding="utf-8-sig"))
    root = local_path(data_dir)
    source = split["source"]
    if len({row["sha256"] for row in source}) != len(source):
        raise ValueError("A content group appears more than once")
    if set(split["assignments"]) != {row["sha256"] for row in source}:
        raise ValueError("Split membership does not match source groups")
    if set(split["assignments"].values()) - {"train", "validation", "test"}:
        raise ValueError("Unknown split partition")
    result = []
    for row in source:
        if split["assignments"][row["sha256"]] != phase:
            continue
        row = dict(row)
        if not re.fullmatch(r"[0-9a-f]{64}", row["sha256"]):
            raise ValueError("Invalid content group identifier")
        relative = row.get("relative_path")
        if relative is None:
            # Relocate the original frozen split without reading its old path.
            filename = PureWindowsPath(row["path"]).name
            if filename in {"", ".", ".."}:
                raise ValueError("Invalid legacy source name")
            relative = f"videos/{filename}"
        path = local_path(root / relative)
        if not path.is_relative_to(root) or not path.is_file():
            raise ValueError("Source path escapes the explicit data directory")
        if type(row["positive"]) is not bool:
            raise ValueError("Labels must be explicit booleans")
        row["path"] = str(path)
        result.append(row)
    if not result or {row["positive"] for row in result} != {False, True}:
        raise ValueError("The selected partition requires both classes")
    if phase == "train" and min(sum(row["positive"] == label for row in result)
        for label in (False, True)) < 4:
        raise ValueError("Four-fold training requires four groups per class")
    return sorted(result, key=lambda row: row["sha256"])


def peak_working_set():
    import sys
    if sys.platform == "win32":
        import ctypes
        from ctypes import wintypes
        class Counters(ctypes.Structure):
            _fields_ = [("cb", wintypes.DWORD), ("PageFaultCount", wintypes.DWORD)] + [
                (name, ctypes.c_size_t) for name in ("PeakWorkingSetSize", "WorkingSetSize",
                "QuotaPeakPagedPoolUsage", "QuotaPagedPoolUsage", "QuotaPeakNonPagedPoolUsage",
                "QuotaNonPagedPoolUsage", "PagefileUsage", "PeakPagefileUsage", "PrivateUsage")]
        kernel = ctypes.windll.kernel32
        kernel.GetCurrentProcess.restype = wintypes.HANDLE
        api = ctypes.windll.psapi.GetProcessMemoryInfo
        api.argtypes = [wintypes.HANDLE, ctypes.c_void_p, wintypes.DWORD]
        api.restype = wintypes.BOOL
        values = Counters()
        values.cb = ctypes.sizeof(values)
        if not api(kernel.GetCurrentProcess(), ctypes.byref(values), values.cb):
            raise OSError("Cannot read process memory usage")
        return int(values.PeakWorkingSetSize)
    import resource
    value = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    return int(value if sys.platform == "darwin" else value * 1024)
