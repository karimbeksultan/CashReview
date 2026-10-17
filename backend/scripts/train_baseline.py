"""Train an honest small-data baseline offline; all outputs belong in private storage.

Exact duplicates share a split. Positive labels identify isolated timestamp neighborhoods;
other parts of positive videos are deliberately not treated as known normal examples.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import time
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from cash_review.engine import FeatureExtractor, MODEL_FILE, MODEL_SHA256, SCHEMA_VERSION, extract_video, score_features, pool_video_features

SEED = 26017


def sha(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(1024 * 1024):
            value.update(chunk)
    return value.hexdigest()


def read_source(data_dir: Path) -> list[dict]:
    from openpyxl import load_workbook

    book = load_workbook(data_dir / "open_labels.xlsx", read_only=True, data_only=True)
    try:
        rows = list(book["labels"].iter_rows(values_only=True))
    finally:
        book.close()
    files = {path.stem: path for path in (data_dir / "videos").glob("*.mov")}
    if len(files) != len(rows) - 1:
        raise ValueError("Source video/label count mismatch; resolve before training")
    groups = {}
    for name, raw_timestamp, raw_label, *_ in rows[1:]:
        if name not in files:
            raise ValueError("Label does not resolve to an exact Unicode filename")
        positive = str(raw_label).strip() == "нарушение"
        if str(raw_label).strip() not in ("нарушение", "без нарушения"):
            raise ValueError("Unknown label; do not coerce it to negative")
        timestamp = None
        if raw_timestamp:
            numbers = [int(value) for value in re.findall(r"\d+", str(raw_timestamp))]
            timestamp = float(numbers[0] * 60 + numbers[1] if len(numbers) == 2 else numbers[0])
        if positive and timestamp is None:
            raise ValueError("Positive row lacks a timestamp")
        path = files[name]
        digest = sha(path)
        if digest not in groups:
            groups[digest] = {"sha256": digest, "path": str(path.resolve()), "positive": positive,
                              "names": [], "timestamps": [], "raw_annotations": []}
        group = groups[digest]
        if group["positive"] != positive:
            raise ValueError("Identical video has conflicting binary labels")
        group["names"].append(name)
        if timestamp is not None and timestamp not in group["timestamps"]:
            group["timestamps"].append(timestamp)
        group["raw_annotations"].append({"name": name, "timestamp": raw_timestamp, "label": raw_label})
    return sorted(groups.values(), key=lambda group: group["sha256"])


def frozen_split(groups: list[dict], directory: Path) -> dict:
    path = directory / "split.json"
    fingerprint = hashlib.sha256(json.dumps([(g["sha256"], g["positive"], sorted(g["timestamps"])) for g in groups]).encode()).hexdigest()
    if path.exists():
        split = json.loads(path.read_text(encoding="utf-8"))
        if split["source_fingerprint"] != fingerprint:
            raise ValueError("Frozen split differs from dataset; choose a new artifact directory instead of overwriting the test set")
        return split
    assignments = {}
    rng = np.random.default_rng(SEED)
    for positive in (False, True):
        members = [g["sha256"] for g in groups if g["positive"] == positive]
        rng.shuffle(members)
        if len(members) < 8:
            raise ValueError("At least eight unique clips per class are needed for this frozen split")
        ntest = max(2, round(len(members) * 0.2))
        nval = max(2, round(len(members) * 0.2))
        for index, digest in enumerate(members):
            assignments[digest] = "test" if index < ntest else "validation" if index < ntest + nval else "train"
    split = {"schema_version": 1, "seed": SEED, "source_fingerprint": fingerprint,
             "method": "SHA-256 exact-content groups, class-stratified fixed 60/20/20 approx split",
             "assignments": assignments, "source": groups,
             "annotation_policy": "Union of supplied timestamps retained for ambiguous exact duplicates; no exhaustive event/type truth"}
    path.write_text(json.dumps(split, ensure_ascii=False, indent=2), encoding="utf-8")
    return split


def load_features(group: dict, directory: Path, extractor: FeatureExtractor) -> dict:
    path = directory / "features" / (group["sha256"] + ".npz")
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        with np.load(path, allow_pickle=False) as data:
            return {key: data[key] for key in data.files}
    started = time.monotonic()
    data = extract_video(Path(group["path"]), extractor, time.monotonic() + 600)
    cached = {"times": data["times"], "features": data["features"],
              "duration_sec": np.array(data["duration_sec"]), "feature_sec": np.array(time.monotonic() - started)}
    np.savez_compressed(path, **cached)
    return cached


def run(args):
    from sklearn.linear_model import LogisticRegression
    from sklearn.preprocessing import StandardScaler
    from sklearn.metrics import f1_score, balanced_accuracy_score, precision_score, confusion_matrix

    directory = args.model_dir.resolve()
    directory.mkdir(parents=True, exist_ok=True)
    if sha(directory / MODEL_FILE) != MODEL_SHA256:
        raise ValueError("Run prepare_model.py first; pinned backbone is missing or corrupted")
    groups = read_source(args.data_dir.resolve())
    split = frozen_split(groups, directory)
    extractor = FeatureExtractor(directory)
    data = {}
    for index, group in enumerate(groups):
        if args.model_kind == "video-gated" and split["assignments"][group["sha256"]] == "test":
            continue  # v2 preparation deliberately does not inspect even cached test features.
        data[group["sha256"]] = load_features(group, directory, extractor)
        print(json.dumps({"feature_group": index + 1, "of": len(groups), "split": split["assignments"][group["sha256"]]}), flush=True)
    x, y, sample_weights = [], [], []
    training = [g for g in groups if split["assignments"][g["sha256"]] == "train"]
    counts = {positive: sum(g["positive"] == positive for g in training) for positive in (False, True)}
    for group in training:
        values = data[group["sha256"]]
        mask = np.ones(len(values["times"]), dtype=bool)
        if group["positive"]:
            mask = np.min(np.abs(values["times"][:, None] - np.asarray(group["timestamps"])[None, :]), axis=1) <= 2.0
        selected = values["features"][mask]
        if len(selected) == 0:
            raise ValueError("Positive timestamp has no sampled neighborhood")
        x.append(selected)
        y.extend([int(group["positive"])] * len(selected))
        sample_weights.extend([1 / (len(selected) * counts[group["positive"]])] * len(selected))
    x = np.concatenate(x)
    # Equal clip contribution and class contribution, irrespective of clip duration.
    sample_weights = np.asarray(sample_weights) * len(y) / 2
    scaler = StandardScaler().fit(x, sample_weight=sample_weights)
    classifier = LogisticRegression(C=0.1, max_iter=2000, solver="lbfgs", random_state=SEED)
    classifier.fit(scaler.transform(x), np.asarray(y), sample_weight=sample_weights)
    weights = {"mean": scaler.mean_.astype(np.float32), "scale": scaler.scale_.astype(np.float32),
               "coef": classifier.coef_[0].astype(np.float32), "intercept": classifier.intercept_.astype(np.float32)}
    validation = [g for g in groups if split["assignments"][g["sha256"]] == "validation"]
    truth = np.asarray([int(g["positive"]) for g in validation])
    max_scores = np.asarray([score_features(data[g["sha256"]]["features"], weights).max() for g in validation])
    candidates = []
    for threshold in np.arange(0.05, 0.951, 0.05):
        prediction = (max_scores >= threshold).astype(int)
        candidates.append((float(f1_score(truth, prediction, zero_division=0)), float(balanced_accuracy_score(truth, prediction)), float(threshold)))
    chosen = max(candidates)  # F1, balanced accuracy, then higher threshold; validation only.
    gate_config = {}
    if args.model_kind == "video-gated":
        # Fixed second model: clip-level supervision is appropriate to binary bag labels.
        # Its feature set and regularization are declared before inspecting v2 test results.
        bag_x = np.stack([pool_video_features(data[g["sha256"]]["features"]) for g in training])
        bag_y = np.asarray([int(g["positive"]) for g in training])
        gate_scaler = StandardScaler().fit(bag_x)
        gate_classifier = LogisticRegression(C=0.1, class_weight="balanced", max_iter=2000, solver="lbfgs", random_state=SEED)
        gate_classifier.fit(gate_scaler.transform(bag_x), bag_y)
        gate_weights = {"mean": gate_scaler.mean_.astype(np.float32), "scale": gate_scaler.scale_.astype(np.float32),
                        "coef": gate_classifier.coef_[0].astype(np.float32), "intercept": gate_classifier.intercept_.astype(np.float32)}
        validation_bags = np.stack([pool_video_features(data[g["sha256"]]["features"]) for g in validation])
        gate_scores = score_features(validation_bags, gate_weights)
        gate_candidates = []
        for gate_threshold in np.arange(0.05, 0.951, 0.05):
            bag_prediction = ((gate_scores >= gate_threshold) & (max_scores >= chosen[2])).astype(int)
            gate_candidates.append((float(balanced_accuracy_score(truth, bag_prediction)),
                                    float(precision_score(truth, bag_prediction, zero_division=0)),
                                    float(f1_score(truth, bag_prediction, zero_division=0)), float(gate_threshold)))
        gate_chosen = max(gate_candidates)
        validation_prediction = ((gate_scores >= gate_chosen[3]) & (max_scores >= chosen[2])).astype(int)
        validation_report = {"model_kind": "video-gated", "training_groups": len(training), "validation_groups": len(validation),
                             "fixed_C": 0.1, "video_feature_policy": "mean,std,max,q90 of frozen1032 features; no duration or identity input",
                             "gate_threshold": round(gate_chosen[3], 6), "validation_balanced_accuracy": gate_chosen[0],
                             "validation_precision": gate_chosen[1], "validation_f1": gate_chosen[2],
                             "validation_confusion_true_rows_predicted_columns": confusion_matrix(truth, validation_prediction, labels=[0, 1]).tolist(),
                             "promotable_by_declared_validation_check": gate_chosen[0] > 0.5,
                             "test_used_for_model_or_threshold_selection": False,
                             "test_status": "Previously exposed by v1; any v2 report on that partition is exploratory"}
        (directory / "validation_v2.json").write_text(json.dumps(validation_report, indent=2), encoding="utf-8")
        print(json.dumps({"video_gate_validation": validation_report}), flush=True)
        if gate_chosen[0] <= 0.5:
            raise ValueError("Video-gated model did not exceed validation chance balanced accuracy; do not promote it")
        weights.update({"gate_" + key: value for key, value in gate_weights.items()})
        gate_config = {"gate_threshold": round(gate_chosen[3], 6), "gate_feature_dim": int(bag_x.shape[1]),
                       "gate_feature_policy": "mean,std,max,q90", "gate_C": 0.1,
                       "gate_threshold_policy": "Max validation balanced accuracy, then precision, then F1, then higher threshold",
                       "validation_balanced_accuracy": gate_chosen[0],
                       "evaluation_status": "exploratory_previously_seen_test",
                       "selection_excludes_test": True}
    np.savez_compressed(directory / "baseline.npz", **weights)
    version = 2 if args.model_kind == "video-gated" else 1
    model_version = f"mobilenetv2-local-review-v{version}-" + sha(directory / "baseline.npz")[:10]
    config = {"schema_version": version, "model_version": model_version, "backbone_sha256": MODEL_SHA256,
              "classifier_sha256": sha(directory / "baseline.npz"), "feature_layer": extractor.feature_layer,
              "feature_dim": int(x.shape[1]), "threshold": round(chosen[2], 6), "sample_interval_sec": 1,
              "seed": SEED, "C": 0.1, "created_utc": datetime.now(timezone.utc).isoformat(),
              "training_groups": len(training), "training_samples": len(y), "validation_groups": len(validation),
              "threshold_policy": "Max validation video F1, then balanced accuracy, then higher threshold; fixed C=0.1",
              "split_fingerprint": split["source_fingerprint"], "competition_ready": False,
              "model_kind": "experimental_visual_review_baseline", "score_calibrated": False, **gate_config}
    (directory / "baseline.json").write_text(json.dumps(config, indent=2), encoding="utf-8")
    print(json.dumps({"trained": True, **config}), flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, required=True, help="Local open/ directory containing labels and videos/")
    parser.add_argument("--model-dir", type=Path, required=True, help="Private ignored artifact directory")
    parser.add_argument("--model-kind", choices=("window", "video-gated"), default="window")
    run(parser.parse_args())
